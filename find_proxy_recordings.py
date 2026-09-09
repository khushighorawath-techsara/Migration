#!/usr/bin/env python3
"""
Finds the PROXY-SIDE recording for rows a human review marked as unresolved.

READ-ONLY. Lists S3, reads Zoom start times, writes one local .xlsx.

THE PROBLEM THIS SOLVES
  A proxy incident involves two Zoom meetings running at the same time:

      1. the interview        candidate + interviewer
      2. the proxy session    candidate + proxy person, feeding answers

  The first pass ranked on candidate + date + time and had no notion that the
  host mattered, so it kept landing on meeting 1. The manual review confirmed
  it: of the rows still unresolved, 64 were hosted by the interviewer and 64 by
  a third person -- only 24 by the proxy.

  So here the proxy stops being a ranking bonus and becomes a FILTER. A session
  is only a candidate if the proxy person hosted it.

WIDER TIME WINDOW, DELIBERATELY
  The first pass demanded a start within 45 minutes because it was matching the
  interview against its own scheduled slot. A proxy session does not start on
  the same minute -- it brackets the interview, often opening earlier and
  running past the end. The default window here is 90 minutes either side.

EVERY UNRESOLVED ROW GETS A REASON
  A blank cell is not an answer. Where nothing is found the output says which
  step failed:

      proxy person not named in the sheet
      that person hosts nothing at all in S3
      they host sessions, but none on that date
      they hosted that date, but not with this candidate
      everything matches except the clock, which is N hours out

  That distinction is the point: "never recorded" and "recorded but
  unconfirmed" need completely different follow-up.

USAGE
  python3 find_proxy_recordings.py --xlsx Untitled_spreadsheet.xlsx
  python3 find_proxy_recordings.py --xlsx in.xlsx --out proxy.xlsx --time-window 120
"""

import argparse
import datetime
import importlib.util
import os
import re
import sys
from collections import defaultdict, Counter

try:
    import openpyxl
except ImportError:
    print("Run: pip install openpyxl boto3", file=sys.stderr)
    sys.exit(1)

HELPER = os.environ.get("MATCHER_PATH", "find_missing_meeting_ids.py")
if not os.path.exists(HELPER):
    print(f"{HELPER} not found -- it is needed for name matching, the S3 index "
          f"and the Zoom lookup.", file=sys.stderr)
    sys.exit(1)

_spec = importlib.util.spec_from_file_location("_matcher", HELPER)
M = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(M)

log = M.log


def review_outcome(value) -> str:
    """Collapse the free-text review note into one state.

    The column holds 29 spelling variants of three ideas -- 'B - YES - I',
    'B YES - I', 'B- YES - I' and so on. Stripping to letters makes them one.
    """
    t = re.sub(r"[^a-z]", "", str(value or "").lower())
    if not t:
        return "BLANK"
    if t.endswith("yesp") or t.endswith("yesip"):
        return "YES-P"          # proxy recording already found -- nothing to do
    if "yes" in t and t.endswith("i"):
        return "YES-I"          # found the interviewer's side, not the proxy's
    if "yes" in t:
        return "YES-?"
    if "no" in t:
        return "NO"
    return "OTHER"


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--xlsx", required=True, help="The reviewed spreadsheet.")
    p.add_argument("--out", default="proxy_recordings.xlsx")
    p.add_argument("--time-window", type=int, default=90,
                   help="Minutes either side of the scheduled slot to accept "
                        "(default 90). Wider than the first pass because a proxy "
                        "session brackets the interview rather than starting with it.")
    p.add_argument("--day-window", type=int, default=1,
                   help="Days either side (default 1), absorbing the UTC/IST "
                        "date-labelling difference.")
    p.add_argument("--no-zoom", action="store_true", help="Skip start-time verification.")
    args = p.parse_args()

    log("=== Proxy-side recording lookup — READ ONLY ===\n")

    # ── source ──────────────────────────────────────────────────────────────
    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [c.value for c in ws[1]]
    H = {n: i for i, n in enumerate(header) if n}
    raw = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    log(f"[1/5] {args.xlsx}: {len(raw)} row(s)")

    C = H.get("Confidence")
    outcomes = Counter(review_outcome(r[C]) for r in raw) if C is not None else Counter()
    for k, v in outcomes.most_common():
        log(f"    {v:4d}  {k}")

    need_idx = [i for i, r in enumerate(raw)
                if C is None or review_outcome(r[C]) not in ("YES-P",)]
    log(f"  -> {len(need_idx)} row(s) need a proxy-side recording\n")

    # ── S3 ──────────────────────────────────────────────────────────────────
    log("[2/5] Indexing S3 ...")
    sessions = M.build_index()

    by_host = defaultdict(list)
    for s in sessions:
        for t in M.name_tokens(s["host"]):
            by_host[t].append(s)
    log(f"  indexed {len(by_host)} host token(s)\n")

    # ── Zoom ────────────────────────────────────────────────────────────────
    tok = None
    if not args.no_zoom:
        log("[3/5] Zoom ...")
        tok = M._zoom_token()
        log("")

    # ── search ──────────────────────────────────────────────────────────────
    log("[4/5] Searching for proxy-hosted sessions ...")
    results, tally = {}, Counter()

    for n, i in enumerate(need_idx, 1):
        r = raw[i]
        proxy = r[H["Proxy Person"]] if "Proxy Person" in H else None
        cand  = r[H["Candidate Name"]]
        rec   = int(r[0]) if isinstance(r[0], (int, float)) else None
        sdate, _ = M.parse_source_date(r[1], rec)
        stimes   = M.parse_source_time(r[2], sdate)

        res = {"mid": "", "conf": "", "why": "", "s3cand": "", "s3date": "",
               "start": "", "gap": "", "alts": ""}

        # step 1 -- is a proxy even named?
        if not proxy or not str(proxy).strip():
            res["why"] = "Proxy Person not named in the sheet"
            tally["no proxy named"] += 1
            results[i] = res; continue

        # step 2 -- does that person host anything at all?
        pool, seen = [], set()
        for t in M.name_tokens(proxy):
            for s in by_host.get(t, ()):
                if id(s) not in seen and M.names_match(proxy, s["host"])[0]:
                    seen.add(id(s)); pool.append(s)
        if not pool:
            res["why"] = f"{proxy} hosts no sessions anywhere in S3"
            tally["proxy hosts nothing"] += 1
            results[i] = res; continue

        # step 3 -- on that date? (either date convention)
        if sdate:
            on_date = [s for s in pool if s["date"] and
                       abs((s["date"] - sdate).days) <= args.day_window]
        else:
            on_date = []
        if not on_date:
            res["why"] = (f"{proxy} hosts {len(pool)} session(s) in S3 but none "
                          f"{'on ' + sdate.isoformat() if sdate else '(no usable source date)'}")
            tally["proxy hosted nothing that date"] += 1
            results[i] = res; continue

        # step 4 -- with this candidate? Not required: a proxy session's folder
        # may be named after the proxy, or "Group", if no external participant
        # was resolved. So candidate agreement upgrades confidence rather than
        # gating the result.
        with_cand = [s for s in on_date if M.names_match(cand, s["candidate"])[0]]

        # step 5 -- rank on the real start time
        def score(s):
            if not stimes:
                return 10 ** 6
            at = None
            if tok:
                at, _ = M.zoom_start_ist(tok, s["meeting_id"])
            if at is None:
                at = s["time_min"]
            if at is None:
                return 10 ** 6 - 1
            return min(min(abs(t - at), 1440 - abs(t - at)) for t, _d, _n in stimes)

        pick_pool = with_cand or on_date
        ranked = sorted(((s, score(s)) for s in pick_pool), key=lambda x: x[1])
        best, gap = ranked[0]

        at, _ad = (M.zoom_start_ist(tok, best["meeting_id"]) if tok else (None, None))
        if at is None:
            at = best["time_min"]

        res.update({
            "mid": best["meeting_id"], "s3cand": best["candidate"],
            "s3date": str(best["date"]) if best["date"] else "",
            "start": f"{at//60:02d}:{at%60:02d}" if at is not None else "",
            "gap": gap if gap < 10 ** 5 else "",
            "alts": "; ".join(f"{s['meeting_id']} ({s['date']})" for s, _g in ranked[1:5]),
        })

        if with_cand and gap <= args.time_window:
            res["conf"] = "P1 — proxy hosted, candidate and time both match"
            tally["P1"] += 1
        elif with_cand:
            res["conf"] = "P2 — proxy hosted with this candidate, time is off"
            res["why"] = (f"nearest proxy session is {gap//60}h{gap%60:02d} from the "
                          f"scheduled slot" if isinstance(gap, int) and gap < 10**5
                          else "no clock available to confirm")
            tally["P2"] += 1
        elif gap <= args.time_window:
            res["conf"] = "P3 — proxy hosted at the right time, candidate folder differs"
            res["why"] = (f"S3 candidate folder is {best['candidate']!r}, not "
                          f"{cand!r} — likely no external participant was resolved")
            tally["P3"] += 1
        else:
            res["conf"] = "P4 — proxy hosted that day, nothing else lines up"
            res["why"] = (f"{proxy} hosted {len(on_date)} session(s) that date but none "
                          f"with this candidate or near the scheduled time")
            tally["P4"] += 1

        results[i] = res
        if n % 25 == 0:
            log(f"    ...{n}/{len(need_idx)}")

    # ── report ──────────────────────────────────────────────────────────────
    log("\n=== Proxy-side result ===")
    for k in ("P1", "P2", "P3", "P4", "no proxy named", "proxy hosts nothing",
              "proxy hosted nothing that date"):
        if tally[k]:
            log(f"  {k:34s} {tally[k]:4d}")
    log(f"\n  {tally['P1']} row(s) have a proxy-hosted session with the candidate "
        f"AND the time confirmed")

    # ── write ───────────────────────────────────────────────────────────────
    out = openpyxl.Workbook(); o = out.active; o.title = "Proxy"
    o.append(list(header) + ["Review outcome", "Proxy Meeting ID", "Proxy confidence",
                             "Proxy S3 candidate", "Proxy S3 date",
                             "Proxy start (IST)", "Proxy time gap (min)",
                             "Why not found", "Proxy alternatives"])
    for i, r in enumerate(raw):
        res = results.get(i)
        oc = review_outcome(r[C]) if C is not None else ""
        if res is None:
            o.append(list(r) + [oc, "", "already found (YES-P)", "", "", "", "", "", ""])
        else:
            o.append(list(r) + [oc, res["mid"], res["conf"], res["s3cand"],
                                res["s3date"], res["start"], res["gap"],
                                res["why"], res["alts"]])
    o.freeze_panes = "A2"
    out.save(args.out)
    log(f"\n[5/5] saved {os.path.abspath(args.out)}")
    log("Nothing in S3, Zoom or Salesforce was modified.")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        log("\nInterrupted."); sys.exit(130)
