"""Split the P3 rows by what the S3 candidate folder actually contains.

P3 means the proxy hosted a session at the right time, but the folder names
someone other than our candidate. That is only reassuring if the folder holds
a PLACEHOLDER. If it holds a different real candidate -- especially one that
appears elsewhere in this same sheet -- then that session belongs to their row,
not this one, and the match is wrong.
"""
import openpyxl, re, sys
from collections import Counter

src = sys.argv[1] if len(sys.argv) > 1 else 'proxy_found.xlsx'
ws = openpyxl.load_workbook(src, data_only=True).active
h = [c.value for c in ws[1]]; I = {n: i for i, n in enumerate(h) if n}
rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]

def norm(s):
    return " ".join(re.sub(r'[^a-z0-9 ]', ' ', str(s or '').replace('_',' ').lower()).split())

# every candidate named anywhere in the sheet
sheet_cands = {norm(r[I['Candidate Name']]) for r in rows if r[I['Candidate Name']]}

buckets, examples = Counter(), {}
for n, r in enumerate(rows, 2):
    conf = str(r[I['Proxy confidence']] or '')
    if not conf.startswith('P3'):
        continue
    folder = r[I['Proxy S3 candidate']]
    proxy  = r[I['Proxy Person']]
    f, p = norm(folder), norm(proxy)

    if f in ('group', 'unknown', 'unknown candidate', '') or 'unknown' in f:
        k = 'placeholder (Group/Unknown)'
    elif p and (f == p or f in p or p in f):
        k = "the PROXY's own name"
    elif f in sheet_cands:
        k = 'a DIFFERENT candidate from this sheet  <-- wrong session'
    else:
        k = 'some other real name'
    buckets[k] += 1
    examples.setdefault(k, []).append((n, r[I['Candidate Name']], folder, proxy))

print(f"P3 rows: {sum(buckets.values())}\n")
for k, v in buckets.most_common():
    print(f"  {v:4d}  {k}")
    for e in examples[k][:4]:
        print(f"          row {e[0]:3d}  sheet={e[1]!r} -> folder={e[2]!r}  (proxy {e[3]!r})")
    print()
